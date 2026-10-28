"""Build a single-column Word version of the paper (paper/CardioSentinel_onecolumn.docx).

Run make_onecolumn.py first: its main.aux / main.bbl give the exact numbers used in the PDF, so the
Word file keeps the same numbering (sections I-V and IV-A..., Tables I-XVIII, Figs. 1-13,
Algorithms 1-2, equations (1)-(11), references [1]-[45]). pandoc cannot draw TikZ or algorithmic
blocks, so Figs. 1 and 3 and both algorithms are rendered with Tectonic and inserted as images;
the PDF figures are converted to PNG. Equations become editable Word equations.

Needs pandoc (or pypandoc_binary), python-docx, Pillow, tectonic and pdftoppm.

    python paper/scripts/make_onecolumn.py && python paper/scripts/make_docx.py
    python paper/scripts/make_docx.py --debug      # write the pandoc source only, no conversion
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
from pathlib import Path

PAPER = Path(__file__).resolve().parents[1]
ONE = PAPER / "onecolumn"
WORK = ONE / "docx_build"
SECTIONS = ["abstract", "introduction", "related_work", "proposed_model", "results", "conclusion"]
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
FONT = "Times New Roman"


# ---------------------------------------------------------------- LaTeX helpers
def arg_end(s: str, i: int) -> int:
    """s[i] == '{'; return the index just past its matching '}'."""
    depth = 0
    for j in range(i, len(s)):
        if s[j] == "{" and s[j - 1] != "\\":
            depth += 1
        elif s[j] == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return j + 1
    raise ValueError(f"unbalanced braces at {i}")


def take_macro(s: str, name: str) -> tuple[str, str | None]:
    """Remove the first \\name{...} from s; return (s without it, its argument)."""
    m = re.search(r"\\" + name + r"\s*\{", s)
    if not m:
        return s, None
    end = arg_end(s, m.end() - 1)
    return s[:m.start()] + s[end:], s[m.end():end - 1]


def unwrap_macro(s: str, name: str, nargs_before: int) -> str:
    """\\name{a}{b}{body} -> body (drops the first nargs_before arguments)."""
    while (m := re.search(r"\\" + name + r"\s*\{", s)):
        i = m.end() - 1
        for _ in range(nargs_before):
            i = arg_end(s, i)
            while s[i] in " \n%":
                i += 1
        end = arg_end(s, i)
        s = s[:m.start()] + s[i + 1:end - 1] + s[end:]
    return s


def inline_inputs(s: str, base: Path) -> str:
    return re.sub(r"\\input\{([^}]*)\}",
                  lambda m: inline_inputs((base / (m.group(1) + ".tex")).read_text(), base), s)


def read_aux(aux: Path) -> tuple[dict[str, str], dict[str, str]]:
    labels, cites = {}, {}
    for line in aux.read_text().splitlines():
        if m := re.match(r"\\newlabel\{([^}]*)\}\{\{(.*?)\}\{", line):
            labels[m.group(1)] = re.sub(r"\\mbox\s*\{", "", m.group(2)).replace("{", "").replace("}", "")
        elif m := re.match(r"\\bibcite\{([^}]*)\}\{([^}]*)\}", line):
            cites[m.group(1)] = m.group(2)
    return labels, cites


# ---------------------------------------------------------------- rendering
def render_png(tex_body: str, preamble: str, out: Path) -> None:
    """Compile a snippet on an empty one-column page and save it as a cropped PNG."""
    from PIL import Image, ImageChops

    d = WORK / "render" / out.stem
    d.mkdir(parents=True, exist_ok=True)
    (d / "s.tex").write_text(preamble + "\\pagestyle{empty}\n\\begin{document}\n" + tex_body + "\n\\end{document}\n")
    subprocess.run(["tectonic", "s.tex"], cwd=d, check=True, capture_output=True)
    subprocess.run(["pdftoppm", "-png", "-r", "250", "-singlefile", "s.pdf", "s"], cwd=d, check=True)
    im = Image.open(d / "s.png").convert("RGB")
    bbox = ImageChops.difference(im, Image.new("RGB", im.size, "white")).getbbox()
    im.crop((bbox[0] - 12, bbox[1] - 12, bbox[2] + 12, bbox[3] + 12)).save(out)


def pdf_to_png(pdf: Path, out: Path) -> None:
    subprocess.run(["pdftoppm", "-png", "-r", "250", "-singlefile", str(pdf), str(out.with_suffix(""))], check=True)


# ---------------------------------------------------------------- conversion
def convert_floats(tex: str, labels: dict[str, str], preamble: str) -> str:
    out, pos = [], 0
    pat = re.compile(r"\\begin\{(figure\*?|table\*?|algorithm)\}(\[[^\]]*\])?")
    while (m := pat.search(tex, pos)):
        env = m.group(1)
        end = tex.index(f"\\end{{{env}}}", m.end())
        body = tex[m.end():end]
        out.append(tex[pos:m.start()])
        pos = end + len(f"\\end{{{env}}}")
        label = re.search(r"\\label\{([^}]*)\}", body).group(1)
        num = labels[label]

        if env == "algorithm":
            png = WORK / f"algorithm{num}.png"
            render_png(f"\\setcounter{{algorithm}}{{{int(num) - 1}}}\n\\begin{{algorithm}}[H]{body}\\end{{algorithm}}",
                       preamble, png)
            out.append(f"\n\n\\begin{{center}}\\includegraphics[width=6.3in]{{{png.name}}}\\end{{center}}\n\n")
            continue

        body, caption = take_macro(body, "caption")
        body = re.sub(r"\\label\{[^}]*\}", "", body)
        wide = env.endswith("*")
        if env.startswith("figure"):
            if "tikzpicture" in body:
                png = WORK / f"fig{num}.png"
                render_png(body, preamble, png)
            else:
                name = re.search(r"\\includegraphics(\[[^\]]*\])?\{([^}]*)\}", body).group(2)
                if Path(name).suffix:                    # already an image (e.g. confusion_combined.jpg)
                    png = WORK / name
                    shutil.copy(PAPER / "figures" / name, png)
                else:
                    png = WORK / f"{name}.png"
                    pdf_to_png(PAPER / "figures" / f"{name}.pdf", png)
            width = "6.3in" if wide else "4.4in"
            out.append(f"\\begin{{figure}}\\centering\\includegraphics[width={width}]{{{png.name}}}"
                       f"\\caption{{Fig. {num}. {caption}}}\\end{{figure}}")
        else:
            body = inline_inputs(body, PAPER)
            body = unwrap_macro(body, "resizebox", 2)
            body = re.sub(r"\\(footnotesize|scriptsize|small|tiny|centering)\b", "", body)
            body = re.sub(r"\\setlength\{\\tabcolsep\}\{[^}]*\}", "", body)
            # pandoc drops a whole tabular whose \multicolumn spec contains @{}
            body = re.sub(r"(\\multicolumn\{\d+\}\{)((?:[^{}]|\{\})*)\}",
                          lambda m: m.group(1) + m.group(2).replace("@{}", "") + "}", body)
            out.append(f"\\begin{{table}}\\caption{{TABLE {num}. {caption}}}{body}\\end{{table}}")
    out.append(tex[pos:])
    return "".join(out)


def number_equations(tex: str) -> str:
    n = 0

    def tag() -> str:
        nonlocal n
        n += 1
        return f"\\qquad\\qquad ({n})"

    def env(m: re.Match) -> str:
        name, body = m.group(1), m.group(2)
        if name == "align":
            # Word has no align: one display equation per line, '&' markers dropped
            lines = [l.replace("&", "") for l in body.split("\\\\") if l.strip()]
            lines = [l.replace("\\nonumber", "") if "\\nonumber" in l else l.rstrip() + tag() for l in lines]
            return "\n".join(f"\\[{l.strip()}\\]" for l in lines)
        if "\\end{split}" in body:
            body = body.replace("\\end{split}", tag() + "\\end{split}")
        else:
            body = body.rstrip() + tag()
        return "\\begin{equation}" + body + "\\end{equation}"

    tex = re.sub(r"\\begin\{(equation|align)\}(.*?)\\end\{\1\}", env, tex, flags=re.S)
    tex = re.sub(r"\\label\{eq:[^}]*\}", "", tex)
    print(f"numbered {n} equations")
    return tex


def number_headings(tex: str) -> str:
    sec = sub = subsub = 0

    def head(m: re.Match) -> str:
        nonlocal sec, sub, subsub
        level, title_start = m.group(1), m.end() - 1
        title = tex_ref[0][title_start + 1:arg_end(tex_ref[0], title_start) - 1]
        if level == "section":
            sec, sub, subsub = sec + 1, 0, 0
            return f"\\section{{{ROMAN[sec - 1]}. {title}}}"
        if level == "subsection":
            sub, subsub = sub + 1, 0
            return f"\\subsection{{{chr(64 + sub)}. {title}}}"
        subsub += 1
        return f"\\subsubsection{{{subsub}) {title}}}"

    out, pos = [], 0
    tex_ref = [tex]
    for m in re.finditer(r"\\(section|subsection|subsubsection)\{", tex):
        if m.start() < pos:
            continue
        out.append(tex[pos:m.start()])
        out.append(head(m))
        pos = arg_end(tex, m.end() - 1)
    out.append(tex[pos:])
    return "".join(out)


def references(bbl: str) -> str:
    items = re.split(r"\\bibitem\{[^}]*\}", bbl)[1:]
    items[-1] = items[-1].replace("\\end{thebibliography}", "")
    # "{[}" keeps pandoc from reading [n] as an optional argument
    paras = [f"{{[}}{i}{{]}}~" + re.sub(r"\s+", " ", it.replace("\\newblock", " ")).strip()
             for i, it in enumerate(items, 1)]
    return "\\section{References}\n\n" + "\n\n".join(paras)


def build_source(labels, cites, preamble) -> str:
    tex = "\n\n".join((PAPER / "sections" / f"{s}.tex").read_text() for s in SECTIONS)
    tex = inline_inputs(tex, PAPER)
    tex = re.sub(r"(?<!\\)%.*", "", tex)                       # comments
    tex = convert_floats(tex, labels, preamble)
    tex = number_equations(tex)
    tex = number_headings(tex)
    tex = re.sub(r"\\begin\{abstract\}", r"\\noindent\\textbf{\\textit{Abstract}}---", tex)
    tex = tex.replace("\\end{abstract}", "")
    tex = re.sub(r"\\begin\{IEEEkeywords\}", r"\\noindent\\textbf{\\textit{Index Terms}}---", tex)
    tex = tex.replace("\\end{IEEEkeywords}", "")
    tex = re.sub(r"\\IEEEPARstart\{(\w)\}\{(\w+)\}", r"\1\2", tex)
    # $^{a}$ has no base, which Word shows as an empty placeholder box
    tex = re.sub(r"\$\^\{?(\w+)\}?\$", r"\\textsuperscript{\1}", tex)
    tex = re.sub(r"\\eqref\{([^}]*)\}", lambda m: f"({labels[m.group(1)]})", tex)
    tex = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels[m.group(1)], tex)
    # {[}n{]}: a bare [n] after \\ in a table row would be read as row spacing
    tex = re.sub(r"\\cite\{([^}]*)\}",
                 lambda m: ", ".join(f"{{[}}{cites[k.strip()]}{{]}}" for k in m.group(1).split(",")), tex)
    missing = re.findall(r"\\(ref|cite|label)\{", tex)
    assert not [x for x in missing if x != "label"], f"unresolved: {missing}"
    tex = re.sub(r"\\label\{[^}]*\}", "", tex)

    main = (PAPER / "main.tex").read_text()
    title = re.search(r"\\title\{(.*?)\}\n", main, re.S).group(1)
    author = re.search(r"\\author\{(.*?)\}\n", main).group(1)
    return ("\\newcommand{\\todo}[1]{}\n"
            f"\\title{{{title}}}\n\\author{{{author}}}\n\\begin{{document}}\n\\maketitle\n\n"
            + tex + "\n\n" + references((ONE / "main.bbl").read_text()) + "\n\\end{document}\n")


# ---------------------------------------------------------------- Word styling
def style_reference(pandoc: str, ref: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Inches

    ref.write_bytes(subprocess.run([pandoc, "--print-default-data-file", "reference.docx"],
                                   check=True, capture_output=True).stdout)
    doc = Document(ref)
    sizes = {"Normal": 11, "Body Text": 11, "First Paragraph": 11, "Compact": 11, "Title": 20, "Author": 12,
             "Heading 1": 12, "Heading 2": 11, "Heading 3": 11, "Caption": 9, "Table Caption": 9,
             "Image Caption": 9}
    for st in doc.styles:
        if st.type != 1 or st.name not in sizes:
            continue
        st.font.name, st.font.size = FONT, Pt(sizes[st.name])
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT)
        st.font.bold = st.name.startswith("Heading") or None
        st.font.italic = True if st.name in ("Heading 2", "Heading 3") else None
        pf = st.paragraph_format
        if st.name in ("Body Text", "First Paragraph"):
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_after = Pt(6)
        if st.name in ("Title", "Author", "Caption", "Table Caption", "Image Caption", "Heading 1"):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if st.name.startswith("Heading"):
            pf.space_before, pf.space_after = Pt(12), Pt(6)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    doc.save(ref)


def finish(docx: Path) -> None:
    """Booktabs-style rules on every table, 9-pt table text, centred tables and images."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document(docx)

    def border(el, side: str, sz: int) -> None:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single"), b.set(qn("w:sz"), str(sz)), b.set(qn("w:color"), "000000")
        el.append(b)

    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblBorders")) + tblPr.findall(qn("w:tblW")):
            tblPr.remove(old)
        width = OxmlElement("w:tblW")                      # full text width
        width.set(qn("w:w"), "5000"), width.set(qn("w:type"), "pct")
        tblPr.append(width)
        # column widths proportional to the longest entry (capped, so text columns wrap)
        for w in t._tbl.xpath(".//w:tcW | .//w:tblLayout"):
            w.getparent().remove(w)
        cols = t._tbl.xpath("./w:tblGrid/w:gridCol")
        need = [7] * len(cols)
        for tr in t._tbl.xpath("./w:tr"):
            tcs = tr.xpath("./w:tc")
            if len(tcs) != len(cols):                          # skip rows with spanning cells
                continue
            for i, tc in enumerate(tcs):
                need[i] = max(need[i], min(len("".join(tc.itertext()).strip()), 34))
        for col, n in zip(cols, need):
            col.set(qn("w:w"), str(round(9360 * n / sum(need))))   # 6.5 in text width, in twips
        borders = OxmlElement("w:tblBorders")
        border(borders, "top", 12), border(borders, "bottom", 12)
        tblPr.append(borders)
        for cell in t.rows[0].cells:
            tcPr = cell._tc.get_or_add_tcPr()
            b = OxmlElement("w:tcBorders")
            border(b, "bottom", 6)
            tcPr.append(b)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(9)
    for p in doc.paragraphs:
        if p._p.xpath(".//pic:pic"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(docx)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--debug", action="store_true", help="write the pandoc source only")
    ap.add_argument("--pandoc", default=shutil.which("pandoc"), help="path to the pandoc binary")
    args = ap.parse_args()
    if not (ONE / "main.aux").exists():
        raise SystemExit("run paper/scripts/make_onecolumn.py first (needs onecolumn/main.aux and main.bbl)")

    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir()
    labels, cites = read_aux(ONE / "main.aux")
    preamble = (ONE / "main.tex").read_text().split("\\begin{document}")[0]
    src = build_source(labels, cites, preamble)
    (WORK / "paper.tex").write_text(src)
    if args.debug:
        print(f"wrote {WORK / 'paper.tex'}")
        return

    if not args.pandoc:
        import pypandoc
        args.pandoc = pypandoc.get_pandoc_path()
    ref = WORK / "reference.docx"
    style_reference(args.pandoc, ref)
    out = PAPER / "CardioSentinel_onecolumn.docx"
    subprocess.run([args.pandoc, "paper.tex", "-f", "latex", "-t", "docx", "--reference-doc", ref.name,
                    "-o", str(out)], cwd=WORK, check=True)
    finish(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
